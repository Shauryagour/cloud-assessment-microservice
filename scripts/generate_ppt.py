"""Generate professional PowerPoint Presentation (.pptx) for Cloud Assessment Project 1."""
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def create_deck():
    prs = Presentation()
    # Set slide dimensions to widescreen 16:9 (13.333 x 7.5 inches)
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_slide_layout = prs.slide_layouts[6] # blank layout

    # Color Palette
    BG_DARK = RGBColor(15, 23, 42)      # #0f172a Deep Slate Navy
    BG_CARD = RGBColor(30, 41, 59)      # #1e293b Slate Card
    TEXT_WHITE = RGBColor(248, 250, 252) # #f8fafc
    TEXT_MUTED = RGBColor(148, 163, 184) # #94a3b8
    ACCENT_ORANGE = RGBColor(255, 153, 0) # #ff9900 AWS Orange
    ACCENT_BLUE = RGBColor(56, 189, 248)  # #38bdf8 Sky Blue
    ACCENT_GREEN = RGBColor(74, 222, 128) # #4ade80 Emerald Green
    BORDER_COLOR = RGBColor(51, 65, 85)   # #334155

    def add_header(slide, title_text, category_text="AWS CLOUD & DISTRIBUTED SYSTEMS"):
        # Header category
        cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.5), Inches(11.7), Inches(0.4))
        tf_cat = cat_box.text_frame
        tf_cat.word_wrap = True
        p_cat = tf_cat.paragraphs[0]
        p_cat.text = category_text.upper()
        p_cat.font.size = Pt(11)
        p_cat.font.bold = True
        p_cat.font.color.rgb = ACCENT_ORANGE
        p_cat.font.name = "Arial"

        # Main Title
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.85), Inches(11.7), Inches(0.8))
        tf = title_box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = title_text
        p.font.size = Pt(24)
        p.font.bold = True
        p.font.color.rgb = TEXT_WHITE
        p.font.name = "Arial"

    def set_slide_background(slide, color=BG_DARK):
        background = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
        background.fill.solid()
        background.fill.fore_color.rgb = color
        background.line.fill.background() # no line
        return background

    # ==========================================
    # SLIDE 1: Title Slide
    # ==========================================
    s1 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(s1)

    # Accent decorative bar
    bar = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.8), Inches(0.8), Inches(0.1))
    bar.fill.solid()
    bar.fill.fore_color.rgb = ACCENT_ORANGE
    bar.line.fill.background()

    # Title & Subtitle Box
    title_box = s1.shapes.add_textbox(Inches(0.8), Inches(2.1), Inches(11.7), Inches(3.2))
    tf1 = title_box.text_frame
    tf1.word_wrap = True

    p1 = tf1.paragraphs[0]
    p1.text = "Containerized Microservice on AWS"
    p1.font.size = Pt(36)
    p1.font.bold = True
    p1.font.color.rgb = TEXT_WHITE
    p1.font.name = "Arial"

    p2 = tf1.add_paragraph()
    p2.text = "Dockerized FastAPI App, ECR Container Registry, EC2 Deployment & GitHub Actions CI/CD"
    p2.font.size = Pt(18)
    p2.font.color.rgb = ACCENT_BLUE
    p2.font.name = "Arial"
    p2.space_before = Pt(14)

    # Meta Info Card at Bottom
    meta_card = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(5.4), Inches(11.7), Inches(1.3))
    meta_card.fill.solid()
    meta_card.fill.fore_color.rgb = BG_CARD
    meta_card.line.color.rgb = BORDER_COLOR

    meta_tf = meta_card.text_frame
    meta_tf.word_wrap = True
    mp1 = meta_tf.paragraphs[0]
    mp1.text = "Presenter: Shaurya Gour    |    Project 1: Cloud & Distributed Systems Assessment"
    mp1.font.size = Pt(14)
    mp1.font.bold = True
    mp1.font.color.rgb = TEXT_WHITE
    mp1.font.name = "Arial"

    mp2 = meta_tf.add_paragraph()
    mp2.text = "Tech Stack: Python 3.11 • FastAPI • Docker • AWS ECR • AWS EC2 • GitHub Actions CI/CD • Pytest"
    mp2.font.size = Pt(12)
    mp2.font.color.rgb = ACCENT_ORANGE
    mp2.font.name = "Arial"
    mp2.space_before = Pt(6)

    # ==========================================
    # SLIDE 2: Objectives & Assessment Scope
    # ==========================================
    s2 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(s2)
    add_header(s2, "Assessment Scope & Architecture Objectives")

    cards_data_s2 = [
        ("1. Production Microservice", "Engineered a high-performance REST microservice using FastAPI with Pydantic validation, structured routing, and OpenAPI/Swagger documentation.", ACCENT_BLUE),
        ("2. Containerization (Docker)", "Encapsulated dependencies into an isolated, multi-stage Python 3.11 slim image with non-root security and healthcheck probes.", ACCENT_ORANGE),
        ("3. AWS Cloud Pipeline", "Configured automated image publishing to Amazon Elastic Container Registry (ECR) and deployment orchestration onto Amazon EC2 virtual machines.", ACCENT_GREEN),
        ("4. Automated CI/CD & Gating", "Implemented dual-layer test gating: Git pre-push hook locally and GitHub Actions remotely, strictly blocking broken code pushes.", ACCENT_BLUE),
    ]

    for i, (title, desc, accent) in enumerate(cards_data_s2):
        col = i % 2
        row = i // 2
        left = Inches(0.8 + col * 5.95)
        top = Inches(1.9 + row * 2.5)

        card = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, Inches(5.75), Inches(2.2))
        card.fill.solid()
        card.fill.fore_color.rgb = BG_CARD
        card.line.color.rgb = BORDER_COLOR

        ctf = card.text_frame
        ctf.word_wrap = True
        cp1 = ctf.paragraphs[0]
        cp1.text = title
        cp1.font.size = Pt(16)
        cp1.font.bold = True
        cp1.font.color.rgb = accent
        cp1.font.name = "Arial"

        cp2 = ctf.add_paragraph()
        cp2.text = desc
        cp2.font.size = Pt(12)
        cp2.font.color.rgb = TEXT_MUTED
        cp2.font.name = "Arial"
        cp2.space_before = Pt(8)

    # ==========================================
    # SLIDE 3: System Architecture (client -> EC2 -> container)
    # ==========================================
    s3 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(s3)
    add_header(s3, "Runtime Architecture: Client → EC2 → Container")

    # Three Flow Columns
    flow_steps = [
        ("Layer 1: Client Ingress", "HTTP Requests (:80)", "Web Browsers, Postman, or Evaluator clients invoke REST APIs via standard HTTP port 80 against the EC2 instance's public IP or DNS.", ACCENT_BLUE),
        ("Layer 2: AWS EC2 Host", "Virtual Machine Layer", "Amazon EC2 (Amazon Linux / Ubuntu) handles traffic routing, firewall security groups, and hosts the active Docker Engine runtime daemon.", ACCENT_ORANGE),
        ("Layer 3: Docker Container", "Isolated Microservice (:8000)", "Docker forwards host Port 80 → Container Port 8000 running Uvicorn ASGI server + FastAPI with in-memory state and health probes.", ACCENT_GREEN)
    ]

    for i, (layer_title, tag, text, accent) in enumerate(flow_steps):
        left = Inches(0.8 + i * 3.97)
        card = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, Inches(1.9), Inches(3.8), Inches(4.8))
        card.fill.solid()
        card.fill.fore_color.rgb = BG_CARD
        card.line.color.rgb = BORDER_COLOR

        ctf = card.text_frame
        ctf.word_wrap = True
        cp1 = ctf.paragraphs[0]
        cp1.text = layer_title
        cp1.font.size = Pt(16)
        cp1.font.bold = True
        cp1.font.color.rgb = accent
        cp1.font.name = "Arial"

        cp_tag = ctf.add_paragraph()
        cp_tag.text = tag
        cp_tag.font.size = Pt(11)
        cp_tag.font.bold = True
        cp_tag.font.color.rgb = TEXT_WHITE
        cp_tag.font.name = "Arial"
        cp_tag.space_before = Pt(4)

        cp_desc = ctf.add_paragraph()
        cp_desc.text = text
        cp_desc.font.size = Pt(12)
        cp_desc.font.color.rgb = TEXT_MUTED
        cp_desc.font.name = "Arial"
        cp_desc.space_before = Pt(12)

    # ==========================================
    # SLIDE 4: Microservice Features & Endpoints
    # ==========================================
    s4 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(s4)
    add_header(s4, "REST Microservice Endpoints & Features")

    endpoints = [
        ("GET /health", "Target Group Health Monitoring", "Returns HTTP 200, status: 'healthy', uptime tracking, and service identifier for AWS ALB / EC2 health checks.", ACCENT_GREEN),
        ("GET /api/info", "Container Runtime Introspection", "Returns runtime OS platform, container hostname, and Python version to prove container isolation.", ACCENT_BLUE),
        ("GET /api/items", "Query & Resource Listing", "Fetches items collection with optional query filtering (?status=pending) and pagination limits.", ACCENT_BLUE),
        ("POST /api/items", "Pydantic Schema Validation", "Validates payload schema; generates unique ID and timestamp. Returns HTTP 201 or HTTP 422 on bad input.", ACCENT_ORANGE),
        ("GET & DELETE /api/items/{id}", "Granular CRUD Operations", "Fetches specific item or removes it with structured HTTP 404 error handling for invalid resource IDs.", ACCENT_BLUE),
        ("GET /docs & /redoc", "Interactive OpenAPI Documentation", "Auto-generated Swagger UI allowing evaluators to test all endpoints interactively in browser.", ACCENT_GREEN),
    ]

    for i, (ep, ep_title, ep_desc, accent) in enumerate(endpoints):
        col = i % 2
        row = i // 2
        left = Inches(0.8 + col * 5.95)
        top = Inches(1.85 + row * 1.68)

        card = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, Inches(5.75), Inches(1.5))
        card.fill.solid()
        card.fill.fore_color.rgb = BG_CARD
        card.line.color.rgb = BORDER_COLOR

        ctf = card.text_frame
        ctf.word_wrap = True
        cp1 = ctf.paragraphs[0]
        cp1.text = f"{ep}  —  {ep_title}"
        cp1.font.size = Pt(13)
        cp1.font.bold = True
        cp1.font.color.rgb = accent
        cp1.font.name = "Arial"

        cp2 = ctf.add_paragraph()
        cp2.text = ep_desc
        cp2.font.size = Pt(11)
        cp2.font.color.rgb = TEXT_MUTED
        cp2.font.name = "Arial"
        cp2.space_before = Pt(4)

    # ==========================================
    # SLIDE 5: Automated Testing & Push Blocking
    # ==========================================
    s5 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(s5)
    add_header(s5, "Automated Test Suite & Push-Blocking Protection")

    # Left box: Test suite
    left_card = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.9), Inches(5.75), Inches(4.8))
    left_card.fill.solid()
    left_card.fill.fore_color.rgb = BG_CARD
    left_card.line.color.rgb = BORDER_COLOR

    ltf = left_card.text_frame
    ltf.word_wrap = True
    lp1 = ltf.paragraphs[0]
    lp1.text = "8 Comprehensive Test Cases (pytest)"
    lp1.font.size = Pt(16)
    lp1.font.bold = True
    lp1.font.color.rgb = ACCENT_GREEN
    lp1.font.name = "Arial"

    tests_list = [
        "1. test_health_check: Verifies status 200 & uptime",
        "2. test_root_endpoint: Validates service metadata & URLs",
        "3. test_system_info_endpoint: Verifies container hostname",
        "4. test_list_items: Checks pre-seeded dataset",
        "5. test_create_item_success: Validates 201 Created & ID",
        "6. test_create_item_validation_error: Validates 422 Unprocessable",
        "7. test_get_item_by_id_and_not_found: Validates 200 & 404",
        "8. test_delete_item_success_and_not_found: Validates deletion"
    ]
    for t in tests_list:
        p = ltf.add_paragraph()
        p.text = t
        p.font.size = Pt(11)
        p.font.color.rgb = TEXT_WHITE
        p.font.name = "Arial"
        p.space_before = Pt(4)

    # Right box: Dual Layer Blocking
    right_card = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.75), Inches(1.9), Inches(5.75), Inches(4.8))
    right_card.fill.solid()
    right_card.fill.fore_color.rgb = BG_CARD
    right_card.line.color.rgb = BORDER_COLOR

    rtf = right_card.text_frame
    rtf.word_wrap = True
    rp1 = rtf.paragraphs[0]
    rp1.text = "Dual-Layer Push Blocking Mechanism"
    rp1.font.size = Pt(16)
    rp1.font.bold = True
    rp1.font.color.rgb = ACCENT_ORANGE
    rp1.font.name = "Arial"

    rp_sub1 = rtf.add_paragraph()
    rp_sub1.text = "Level 1: Git Pre-Push Hook (.githooks/pre-push)"
    rp_sub1.font.size = Pt(13)
    rp_sub1.font.bold = True
    rp_sub1.font.color.rgb = ACCENT_BLUE
    rp_sub1.font.name = "Arial"
    rp_sub1.space_before = Pt(10)

    rp_desc1 = rtf.add_paragraph()
    rp_desc1.text = "• Intercepts every 'git push' command locally.\n• Automatically runs pytest in 0.3s.\n• If ANY test fails, exit code 1 aborts the push immediately before code leaves developer machine."
    rp_desc1.font.size = Pt(11)
    rp_desc1.font.color.rgb = TEXT_MUTED
    rp_desc1.font.name = "Arial"
    rp_desc1.space_before = Pt(4)

    rp_sub2 = rtf.add_paragraph()
    rp_sub2.text = "Level 2: GitHub Actions CI Gate"
    rp_sub2.font.size = Pt(13)
    rp_sub2.font.bold = True
    rp_sub2.font.color.rgb = ACCENT_BLUE
    rp_sub2.font.name = "Arial"
    rp_sub2.space_before = Pt(12)

    rp_desc2 = rtf.add_paragraph()
    rp_desc2.text = "• Runs test job in Ubuntu runner on every commit.\n• Downstream Docker build & AWS deployment jobs strictly depend on 'needs: test'.\n• Broken builds are halted automatically."
    rp_desc2.font.size = Pt(11)
    rp_desc2.font.color.rgb = TEXT_MUTED
    rp_desc2.font.name = "Arial"
    rp_desc2.space_before = Pt(4)

    # ==========================================
    # SLIDE 6: CI/CD & AWS Deployment Workflow
    # ==========================================
    s6 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(s6)
    add_header(s6, "Automated GitHub Actions CI/CD Pipeline")

    pipeline_steps = [
        ("1. Git Commit & Push", "Developer triggers push", "Local pre-push hook validates all 8 test cases before uploading to GitHub repository.", ACCENT_BLUE),
        ("2. Test Suite Execution", "GitHub Actions Runner", "Runs pytest in clean Python 3.11 environment on ubuntu-latest container runner.", ACCENT_GREEN),
        ("3. Docker Build & Tag", "Local Verification", "Builds multi-stage Docker image, tags with commit SHA + latest for ECR registry.", ACCENT_ORANGE),
        ("4. AWS ECR Publishing", "Amazon Container Registry", "Authenticates with AWS and pushes production image to Amazon ECR repository.", ACCENT_ORANGE),
        ("5. EC2 Container Deploy", "SSH / Remote Daemon", "Connects to EC2 instance, pulls latest ECR image, restarts container on port 80.", ACCENT_GREEN)
    ]

    for i, (title, subtitle, desc, accent) in enumerate(pipeline_steps):
        left = Inches(0.8 + i * 2.38)
        card = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, Inches(1.9), Inches(2.25), Inches(4.8))
        card.fill.solid()
        card.fill.fore_color.rgb = BG_CARD
        card.line.color.rgb = BORDER_COLOR

        ctf = card.text_frame
        ctf.word_wrap = True
        cp1 = ctf.paragraphs[0]
        cp1.text = title
        cp1.font.size = Pt(13)
        cp1.font.bold = True
        cp1.font.color.rgb = accent
        cp1.font.name = "Arial"

        cp_sub = ctf.add_paragraph()
        cp_sub.text = subtitle
        cp_sub.font.size = Pt(10)
        cp_sub.font.bold = True
        cp_sub.font.color.rgb = TEXT_WHITE
        cp_sub.font.name = "Arial"
        cp_sub.space_before = Pt(4)

        cp_desc = ctf.add_paragraph()
        cp_desc.text = desc
        cp_desc.font.size = Pt(11)
        cp_desc.font.color.rgb = TEXT_MUTED
        cp_desc.font.name = "Arial"
        cp_desc.space_before = Pt(10)

    # ==========================================
    # SLIDE 7: Summary & Live Verification
    # ==========================================
    s7 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(s7)
    add_header(s7, "Project Results & Assessment Summary")

    # Big Result Card
    summary_card = s7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.9), Inches(11.7), Inches(4.8))
    summary_card.fill.solid()
    summary_card.fill.fore_color.rgb = BG_CARD
    summary_card.line.color.rgb = BORDER_COLOR

    stf = summary_card.text_frame
    stf.word_wrap = True

    sp1 = stf.paragraphs[0]
    sp1.text = "All Rubric Requirements Satisfied & Verified"
    sp1.font.size = Pt(18)
    sp1.font.bold = True
    sp1.font.color.rgb = ACCENT_GREEN
    sp1.font.name = "Arial"

    points = [
        ("Public GitHub Repository", "https://github.com/Shauryagour/cloud-assessment-microservice (Clean & Public)"),
        ("GitHub Actions CI/CD Pipeline", "Active & 100% Green (Conclusion: success across all test & build stages)"),
        ("Automated Test Suite", "8 / 8 Tests Passing in 0.3s (Health checks, Pydantic validation, CRUD, 404/422 handling)"),
        ("Push-Blocking Security", "Active Git Pre-Push Hook + GitHub Actions deployment gate prevent broken code"),
        ("Architecture Documentation", "Comprehensive README with Mermaid / ASCII diagrams and step-by-step reproduction"),
        ("Containerization & Cloud Ready", "Multi-stage Dockerfile, docker-compose.yml, ECR push workflow, and EC2 deploy script")
    ]

    for label, val in points:
        p = stf.add_paragraph()
        p.text = f"✔  {label}: {val}"
        p.font.size = Pt(12)
        p.font.color.rgb = TEXT_WHITE
        p.font.name = "Arial"
        p.space_before = Pt(8)

    output_path1 = "/Users/Shaurya/Documents/Cloud_assessment/cloud_assessment_presentation.pptx"
    output_path2 = "/Users/Shaurya/Documents/Cloud_assessment/presentation.pptx"
    prs.save(output_path1)
    prs.save(output_path2)
    print(f"Presentation saved successfully to:\n- {output_path1}\n- {output_path2}")

if __name__ == "__main__":
    create_deck()
